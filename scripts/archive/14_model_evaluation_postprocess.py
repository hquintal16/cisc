#!/usr/bin/env python3
"""Post-process completed county-level WWA evaluation outputs.

This script reads the completed Stage 13 metric Parquet files and performs all
remaining downstream model-evaluation work for the full 2002-2023 window only.

Input
-----
data/out/13_performance/metrics_county/
    metrics_heatindex_all_heat_multiwindow.parquet
    metrics_precipitation_all_precip_multiwindow.parquet

Filter
------
window_tag == "20020101_20231231"

Outputs
-------
1. confusion_matrices.csv / .parquet
2. summary_f1_heatindex_20020101_20231231_all_rps.csv / .parquet
3. summary_f1_precipitation_20020101_20231231_all_rps.csv / .parquet
4. f1_comparison_vs_native_heatindex_20020101_20231231_all_rps.csv / .parquet
5. f1_comparison_vs_native_precipitation_20020101_20231231_all_rps.csv / .parquet
6. performance_metric_tables/*.xlsx
"""

from __future__ import annotations

import os
import tempfile
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from scipy import stats


ROOT = Path(
    os.environ.get(
        "CISC_ROOT",
        "/proj/sebastianlab/users/hquintal/cisc",
    )
)

DATA_OUT = Path(
    os.environ.get(
        "CISC_DATA_OUT",
        str(ROOT / "data/out"),
    )
)

PERF_BASE = Path(
    os.environ.get(
        "CISC_PERF_DIR",
        str(DATA_OUT / "13_performance"),
    )
)

METRICS_DIR = PERF_BASE / "metrics_county"
TABLE_DIR = PERF_BASE / "performance_metric_tables"

HEAT_FILE = METRICS_DIR / "metrics_heatindex_all_heat_multiwindow.parquet"
PRECIP_FILE = (
    METRICS_DIR
    / "metrics_precipitation_all_precip_multiwindow.parquet"
)

WINDOW_TAG = "20020101_20231231"

RETURN_PERIODS = (5, 10, 25, 50)

SCENARIO_BY_RP = {
    5: "filtered_rp5",
    10: "filtered_rp10",
    25: "filtered_rp25",
    50: "filtered_rp50",
}

RP_BY_SCENARIO = {
    scenario: return_period
    for return_period, scenario in SCENARIO_BY_RP.items()
}

RESOLUTION_ORDER = (
    "Native",
    "STM1 bilinear",
    "STM1 max",
    "STM2 bilinear",
    "STM2 max",
)

DATASET_LABELS = {
    ("heatindex", "heatindex_0.25"): "Native",
    ("heatindex", "heatindex_0.30_bilinear"): "STM1 bilinear",
    ("heatindex", "heatindex_0.30_max"): "STM1 max",
    ("heatindex", "heatindex_0.36_bilinear"): "STM2 bilinear",
    ("heatindex", "heatindex_0.36_max"): "STM2 max",
    ("precipitation", "precipitation_0.25"): "Native",
    ("precipitation", "precipitation_0.11_bilinear"): "STM1 bilinear",
    ("precipitation", "precipitation_0.11_max"): "STM1 max",
    ("precipitation", "precipitation_0.14_bilinear"): "STM2 bilinear",
    ("precipitation", "precipitation_0.14_max"): "STM2 max",
}

NATIVE_DATASET = {
    "heatindex": "heatindex_0.25",
    "precipitation": "precipitation_0.25",
}

METRIC_NAMES = (
    "precision",
    "recall",
    "specificity",
    "accuracy",
    "f1",
    "f2",
)

COUNT_NAMES = (
    "TP",
    "FP",
    "FN",
    "TN",
    "total",
)


def log(message: str) -> None:
    print(
        f"[{datetime.now():%Y-%m-%d %H:%M:%S}] {message}",
        flush=True,
    )


def atomic_csv(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp.csv")
    frame.to_csv(temporary, index=False)
    os.replace(temporary, path)


def atomic_parquet(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)

    fd, temporary_name = tempfile.mkstemp(
        prefix=f"{path.stem}.",
        suffix=".tmp.parquet",
        dir=path.parent,
    )
    os.close(fd)
    temporary = Path(temporary_name)

    try:
        frame.to_parquet(
            temporary,
            index=False,
            engine="pyarrow",
            compression="zstd",
        )
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def write_both(frame: pd.DataFrame, csv_path: Path) -> None:
    atomic_csv(frame, csv_path)
    atomic_parquet(frame, csv_path.with_suffix(".parquet"))


def read_metrics(path: Path) -> pd.DataFrame:
    if not path.exists():
        raise FileNotFoundError(f"Missing metrics file: {path}")

    frame = pd.read_parquet(path)

    required = {
        "scenario",
        "dataset_tag",
        "window_tag",
        "GEOID",
        "TP",
        "FP",
        "FN",
        "TN",
        "f1",
        "var_group",
    }

    missing = required.difference(frame.columns)

    if missing:
        raise ValueError(
            f"{path.name} is missing columns: "
            + ", ".join(sorted(missing))
        )

    frame = frame.copy()
    frame["scenario"] = frame["scenario"].astype("string")
    frame["dataset_tag"] = frame["dataset_tag"].astype("string")
    frame["window_tag"] = frame["window_tag"].astype("string")
    frame["GEOID"] = frame["GEOID"].astype("string").str.zfill(5)
    frame["var_group"] = (
        frame["var_group"]
        .astype("string")
        .str.lower()
        .str.strip()
        .replace(
            {
                "heat index": "heatindex",
                "heat_index": "heatindex",
                "precip": "precipitation",
            }
        )
    )

    for column in ("TP", "FP", "FN", "TN", "f1"):
        frame[column] = pd.to_numeric(
            frame[column],
            errors="coerce",
        )

    frame = frame[
        frame["window_tag"].eq(WINDOW_TAG)
        & frame["scenario"].isin(RP_BY_SCENARIO)
        & frame["var_group"].isin(
            {"heatindex", "precipitation"}
        )
    ].copy()

    frame["return_period"] = (
        frame["scenario"]
        .map(RP_BY_SCENARIO)
        .astype("Int64")
    )

    return frame


def derive_metrics(frame: pd.DataFrame) -> pd.DataFrame:
    output = frame.copy()

    output["total"] = (
        output["TP"]
        + output["FP"]
        + output["FN"]
        + output["TN"]
    )

    def safe_divide(
        numerator: pd.Series,
        denominator: pd.Series,
    ) -> pd.Series:
        return numerator / denominator.replace(0, np.nan)

    output["precision"] = safe_divide(
        output["TP"],
        output["TP"] + output["FP"],
    )
    output["recall"] = safe_divide(
        output["TP"],
        output["TP"] + output["FN"],
    )
    output["specificity"] = safe_divide(
        output["TN"],
        output["TN"] + output["FP"],
    )
    output["accuracy"] = safe_divide(
        output["TP"] + output["TN"],
        output["total"],
    )
    output["f1"] = safe_divide(
        2 * output["TP"],
        2 * output["TP"] + output["FP"] + output["FN"],
    )
    output["f2"] = safe_divide(
        5 * output["TP"],
        5 * output["TP"]
        + 4 * output["FN"]
        + output["FP"],
    )

    return output


def build_confusion_summary(
    metrics: pd.DataFrame,
) -> pd.DataFrame:
    grouped = (
        metrics.groupby(
            [
                "var_group",
                "dataset_tag",
                "return_period",
            ],
            dropna=False,
        )
        .agg(
            TP=("TP", "sum"),
            FP=("FP", "sum"),
            FN=("FN", "sum"),
            TN=("TN", "sum"),
            N_rows=("GEOID", "size"),
        )
        .reset_index()
        .rename(columns={"var_group": "variable"})
    )

    grouped = derive_metrics(grouped)

    grouped["resolution"] = [
        DATASET_LABELS.get(
            (variable, dataset_tag),
            pd.NA,
        )
        for variable, dataset_tag in zip(
            grouped["variable"],
            grouped["dataset_tag"],
        )
    ]

    unknown = grouped[grouped["resolution"].isna()]

    if not unknown.empty:
        log(
            "Excluding unrecognized datasets: "
            + ", ".join(
                f"{row.variable}:{row.dataset_tag}"
                for row in unknown.itertuples(index=False)
            )
        )

    grouped = grouped[
        grouped["resolution"].notna()
    ].copy()

    grouped["_resolution_order"] = (
        grouped["resolution"]
        .map(
            {
                label: index
                for index, label in enumerate(
                    RESOLUTION_ORDER
                )
            }
        )
    )

    grouped.sort_values(
        [
            "variable",
            "_resolution_order",
            "return_period",
        ],
        inplace=True,
    )

    grouped.drop(
        columns="_resolution_order",
        inplace=True,
    )

    grouped.reset_index(
        drop=True,
        inplace=True,
    )

    return grouped


def summarize_county_f1(
    metrics: pd.DataFrame,
    variable: str,
) -> pd.DataFrame:
    subset = metrics[
        metrics["var_group"].eq(variable)
    ].copy()

    output = (
        subset.groupby(
            [
                "scenario",
                "dataset_tag",
                "return_period",
                "window_tag",
            ],
            dropna=False,
        )["f1"]
        .agg(
            n_counties="count",
            mean_f1="mean",
            median_f1="median",
            sd_f1="std",
            min_f1="min",
            max_f1="max",
        )
        .reset_index()
    )

    output.insert(0, "variable", variable)

    output.sort_values(
        [
            "return_period",
            "dataset_tag",
        ],
        inplace=True,
    )

    output.reset_index(
        drop=True,
        inplace=True,
    )

    return output


def format_p(value: float) -> str:
    if not np.isfinite(value):
        return "p = NA"
    if value < 0.001:
        return "p < 0.001"
    return f"p = {value:.3f}"


def paired_test_results(
    dataset_values: np.ndarray,
    native_values: np.ndarray,
) -> dict[str, float]:
    valid = (
        np.isfinite(dataset_values)
        & np.isfinite(native_values)
    )

    dataset_values = dataset_values[valid]
    native_values = native_values[valid]
    differences = dataset_values - native_values

    if differences.size == 0:
        return {
            "diff_mean": np.nan,
            "ci_low": np.nan,
            "ci_high": np.nan,
            "t_p_value": np.nan,
            "wilcox_p_value": np.nan,
        }

    diff_mean = float(np.mean(differences))

    if differences.size >= 2:
        t_result = stats.ttest_rel(
            dataset_values,
            native_values,
            nan_policy="omit",
        )

        t_p_value = float(t_result.pvalue)

        standard_error = stats.sem(
            differences,
            nan_policy="omit",
        )

        if np.isfinite(standard_error):
            if standard_error == 0:
                ci_low = diff_mean
                ci_high = diff_mean
            else:
                critical = stats.t.ppf(
                    0.975,
                    df=differences.size - 1,
                )
                margin = float(
                    critical * standard_error
                )
                ci_low = diff_mean - margin
                ci_high = diff_mean + margin
        else:
            ci_low = np.nan
            ci_high = np.nan
    else:
        t_p_value = np.nan
        ci_low = np.nan
        ci_high = np.nan

    if (
        differences.size >= 1
        and np.any(differences != 0)
    ):
        try:
            wilcox_p_value = float(
                stats.wilcoxon(
                    dataset_values,
                    native_values,
                    zero_method="wilcox",
                    correction=False,
                    alternative="two-sided",
                ).pvalue
            )
        except ValueError:
            wilcox_p_value = np.nan
    else:
        wilcox_p_value = np.nan

    return {
        "diff_mean": diff_mean,
        "ci_low": ci_low,
        "ci_high": ci_high,
        "t_p_value": t_p_value,
        "wilcox_p_value": wilcox_p_value,
    }


def compare_f1_to_native(
    metrics: pd.DataFrame,
    variable: str,
) -> pd.DataFrame:
    subset = metrics[
        metrics["var_group"].eq(variable)
    ].copy()

    native_tag = NATIVE_DATASET[variable]
    rows: list[dict] = []

    for return_period in RETURN_PERIODS:
        scenario = SCENARIO_BY_RP[
            return_period
        ]

        scenario_data = subset[
            subset["scenario"].eq(scenario)
        ].copy()

        native = (
            scenario_data.loc[
                scenario_data["dataset_tag"].eq(native_tag)
                & scenario_data["f1"].notna(),
                ["GEOID", "f1"],
            ]
            .rename(columns={"f1": "native_f1"})
            .drop_duplicates("GEOID")
        )

        other_tags = sorted(
            set(
                scenario_data[
                    "dataset_tag"
                ].dropna().astype(str)
            )
            - {native_tag}
        )

        for dataset_tag in other_tags:
            dataset = (
                scenario_data.loc[
                    scenario_data["dataset_tag"].eq(dataset_tag)
                    & scenario_data["f1"].notna(),
                    ["GEOID", "f1"],
                ]
                .rename(
                    columns={"f1": "dataset_f1"}
                )
                .drop_duplicates("GEOID")
            )

            paired = native.merge(
                dataset,
                on="GEOID",
                how="inner",
            )

            if paired.empty:
                test_values = {
                    "diff_mean": np.nan,
                    "ci_low": np.nan,
                    "ci_high": np.nan,
                    "t_p_value": np.nan,
                    "wilcox_p_value": np.nan,
                }
                mean_dataset = np.nan
                median_dataset = np.nan
                mean_native = np.nan
                median_native = np.nan
            else:
                dataset_values = paired[
                    "dataset_f1"
                ].to_numpy(dtype=float)
                native_values = paired[
                    "native_f1"
                ].to_numpy(dtype=float)

                test_values = paired_test_results(
                    dataset_values,
                    native_values,
                )

                mean_dataset = float(
                    np.mean(dataset_values)
                )
                median_dataset = float(
                    np.median(dataset_values)
                )
                mean_native = float(
                    np.mean(native_values)
                )
                median_native = float(
                    np.median(native_values)
                )

            resolution = DATASET_LABELS.get(
                (variable, dataset_tag),
                dataset_tag,
            )

            rows.append(
                {
                    "variable": variable,
                    "return_period": return_period,
                    "scenario": scenario,
                    "window_tag": WINDOW_TAG,
                    "dataset_tag": dataset_tag,
                    "label": resolution,
                    "mean_dataset": mean_dataset,
                    "median_dataset": median_dataset,
                    "mean_native": mean_native,
                    "median_native": median_native,
                    **test_values,
                    "n_dataset": int(
                        len(dataset)
                    ),
                    "n_native": int(
                        len(native)
                    ),
                    "n_paired": int(
                        len(paired)
                    ),
                    "p_label": format_p(
                        test_values[
                            "t_p_value"
                        ]
                    ),
                }
            )

    output = pd.DataFrame(rows)

    if not output.empty:
        output["_resolution_order"] = (
            output["label"]
            .map(
                {
                    label: index
                    for index, label in enumerate(
                        RESOLUTION_ORDER
                    )
                }
            )
            .fillna(999)
        )

        output.sort_values(
            [
                "return_period",
                "_resolution_order",
                "dataset_tag",
            ],
            inplace=True,
        )

        output.drop(
            columns="_resolution_order",
            inplace=True,
        )

        output.reset_index(
            drop=True,
            inplace=True,
        )

    return output


def metric_matrix(
    confusion: pd.DataFrame,
    variable: str,
    metric: str,
) -> pd.DataFrame:
    subset = confusion[
        confusion["variable"].eq(variable)
    ][
        [
            "resolution",
            "return_period",
            metric,
        ]
    ].copy()

    template = pd.MultiIndex.from_product(
        [
            RESOLUTION_ORDER,
            RETURN_PERIODS,
        ],
        names=[
            "resolution",
            "return_period",
        ],
    ).to_frame(index=False)

    subset = template.merge(
        subset,
        on=[
            "resolution",
            "return_period",
        ],
        how="left",
    )

    wide = subset.pivot(
        index="resolution",
        columns="return_period",
        values=metric,
    )

    wide = wide.reindex(
        index=RESOLUTION_ORDER,
        columns=RETURN_PERIODS,
    )

    wide.columns = [
        f"{return_period}-year RP"
        for return_period in wide.columns
    ]

    wide.reset_index(
        inplace=True,
    )

    return wide


def style_sheet(
    worksheet,
    title: str,
    table: pd.DataFrame,
    metric: bool,
) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet["A1"] = title
    worksheet["A1"].font = Font(
        name="Calibri",
        size=12,
        bold=True,
    )

    header_row = 3
    thin_bottom = Side(
        style="thin",
        color="000000",
    )

    for column_index, column_name in enumerate(
        table.columns,
        start=1,
    ):
        cell = worksheet.cell(
            row=header_row,
            column=column_index,
            value=column_name,
        )
        cell.font = Font(
            name="Calibri",
            size=10,
            bold=True,
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        cell.border = Border(
            bottom=thin_bottom
        )

    for row_index, values in enumerate(
        table.itertuples(
            index=False,
            name=None,
        ),
        start=header_row + 1,
    ):
        for column_index, value in enumerate(
            values,
            start=1,
        ):
            cell = worksheet.cell(
                row=row_index,
                column=column_index,
                value=(
                    None
                    if pd.isna(value)
                    else value
                ),
            )

            cell.font = Font(
                name="Calibri",
                size=10,
                bold=column_index == 1,
            )

            cell.alignment = Alignment(
                horizontal=(
                    "left"
                    if column_index == 1
                    else "center"
                ),
                vertical="center",
            )

            if (
                column_index > 1
                and value is not None
                and not pd.isna(value)
            ):
                cell.number_format = (
                    "0.0000"
                    if metric
                    else "#,##0"
                )

    worksheet.column_dimensions[
        "A"
    ].width = 19

    for column_index in range(
        2,
        len(table.columns) + 1,
    ):
        worksheet.column_dimensions[
            get_column_letter(
                column_index
            )
        ].width = 14

    worksheet.freeze_panes = "B4"
    worksheet.auto_filter.ref = (
        f"A3:"
        f"{get_column_letter(len(table.columns))}"
        f"{3 + len(table)}"
    )


def write_metric_workbook(
    confusion: pd.DataFrame,
    metric: str,
) -> None:
    workbook = Workbook()
    workbook.remove(
        workbook.active
    )

    for variable in (
        "heatindex",
        "precipitation",
    ):
        worksheet = workbook.create_sheet(
            title=variable
        )

        table = metric_matrix(
            confusion,
            variable,
            metric,
        )

        variable_title = (
            "Heat index"
            if variable == "heatindex"
            else "Precipitation"
        )

        style_sheet(
            worksheet,
            (
                f"{variable_title}: {metric} "
                "by grid resolution and "
                "return period"
            ),
            table,
            metric=True,
        )

    output_path = (
        TABLE_DIR
        / f"{metric}.xlsx"
    )

    workbook.save(
        output_path
    )

    log(f"Wrote: {output_path}")


def write_count_workbook(
    confusion: pd.DataFrame,
) -> None:
    workbook = Workbook()
    workbook.remove(
        workbook.active
    )

    for variable in (
        "heatindex",
        "precipitation",
    ):
        for count_name in COUNT_NAMES:
            prefix = (
                "heat"
                if variable == "heatindex"
                else "precip"
            )

            worksheet = workbook.create_sheet(
                title=f"{prefix}_{count_name}"
            )

            table = metric_matrix(
                confusion,
                variable,
                count_name,
            )

            variable_title = (
                "Heat index"
                if variable == "heatindex"
                else "Precipitation"
            )

            style_sheet(
                worksheet,
                (
                    f"{variable_title}: "
                    f"{count_name} by grid "
                    "resolution and return "
                    "period"
                ),
                table,
                metric=False,
            )

    output_path = (
        TABLE_DIR
        / "confusion_matrix_counts.xlsx"
    )

    workbook.save(
        output_path
    )

    log(f"Wrote: {output_path}")


def write_audit_workbook(
    confusion: pd.DataFrame,
) -> None:
    workbook = Workbook()
    workbook.remove(
        workbook.active
    )

    columns = [
        "variable",
        "dataset_tag",
        "resolution",
        "return_period",
        "TP",
        "FP",
        "FN",
        "TN",
        "total",
        "precision",
        "recall",
        "specificity",
        "accuracy",
        "f1",
        "f2",
    ]

    for variable in (
        "heatindex",
        "precipitation",
    ):
        worksheet = workbook.create_sheet(
            title=variable
        )

        subset = confusion[
            confusion["variable"].eq(variable)
        ][columns].copy()

        subset["_resolution_order"] = (
            subset["resolution"].map(
                {
                    label: index
                    for index, label in enumerate(
                        RESOLUTION_ORDER
                    )
                }
            )
        )

        subset.sort_values(
            [
                "return_period",
                "_resolution_order",
            ],
            inplace=True,
        )

        subset.drop(
            columns="_resolution_order",
            inplace=True,
        )

        for column_index, column_name in enumerate(
            subset.columns,
            start=1,
        ):
            cell = worksheet.cell(
                row=1,
                column=column_index,
                value=column_name,
            )
            cell.font = Font(
                name="Calibri",
                size=10,
                bold=True,
            )
            cell.alignment = Alignment(
                horizontal="center"
            )

        for row_index, values in enumerate(
            subset.itertuples(
                index=False,
                name=None,
            ),
            start=2,
        ):
            for column_index, value in enumerate(
                values,
                start=1,
            ):
                cell = worksheet.cell(
                    row=row_index,
                    column=column_index,
                    value=(
                        None
                        if pd.isna(value)
                        else value
                    ),
                )

                if (
                    subset.columns[
                        column_index - 1
                    ]
                    in METRIC_NAMES
                    and value is not None
                    and not pd.isna(value)
                ):
                    cell.number_format = "0.0000"

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = (
            f"A1:"
            f"{get_column_letter(len(subset.columns))}"
            f"{len(subset) + 1}"
        )

        for column_index, column_name in enumerate(
            subset.columns,
            start=1,
        ):
            width = max(
                len(column_name) + 2,
                12,
            )

            worksheet.column_dimensions[
                get_column_letter(column_index)
            ].width = width

    output_path = (
        TABLE_DIR
        / "all_performance_metrics.xlsx"
    )

    workbook.save(
        output_path
    )

    log(f"Wrote: {output_path}")


def main() -> None:
    TABLE_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    log(f"Reading: {HEAT_FILE}")
    heat = read_metrics(
        HEAT_FILE
    )

    log(f"Reading: {PRECIP_FILE}")
    precipitation = read_metrics(
        PRECIP_FILE
    )

    metrics = pd.concat(
        [
            heat,
            precipitation,
        ],
        ignore_index=True,
    )

    expected = {
        (
            variable,
            scenario,
            dataset_tag,
        )
        for variable in (
            "heatindex",
            "precipitation",
        )
        for scenario in SCENARIO_BY_RP.values()
        for label_variable, dataset_tag in DATASET_LABELS
        if label_variable == variable
    }

    observed = set(
        metrics[
            [
                "var_group",
                "scenario",
                "dataset_tag",
            ]
        ]
        .drop_duplicates()
        .itertuples(
            index=False,
            name=None,
        )
    )

    missing = expected.difference(observed)

    if missing:
        raise RuntimeError(
            "Missing Stage 13 variable/scenario/dataset "
            "combinations:\n"
            + "\n".join(
                ":".join(values)
                for values in sorted(missing)
            )
        )

    log(
        f"Rows retained for {WINDOW_TAG}: "
        f"{len(metrics)}"
    )

    confusion = build_confusion_summary(
        metrics
    )

    expected_confusion_rows = (
        2
        * len(RESOLUTION_ORDER)
        * len(RETURN_PERIODS)
    )

    if len(confusion) != expected_confusion_rows:
        raise RuntimeError(
            f"Expected {expected_confusion_rows} "
            "confusion-matrix rows, but generated "
            f"{len(confusion)}"
        )

    confusion_csv = (
        PERF_BASE
        / "confusion_matrices.csv"
    )

    write_both(
        confusion,
        confusion_csv,
    )

    heat_summary = summarize_county_f1(
        metrics,
        "heatindex",
    )

    precipitation_summary = summarize_county_f1(
        metrics,
        "precipitation",
    )

    heat_comparison = compare_f1_to_native(
        metrics,
        "heatindex",
    )

    precipitation_comparison = compare_f1_to_native(
        metrics,
        "precipitation",
    )

    write_both(
        heat_summary,
        METRICS_DIR
        / (
            "summary_f1_heatindex_"
            f"{WINDOW_TAG}_all_rps.csv"
        ),
    )

    write_both(
        precipitation_summary,
        METRICS_DIR
        / (
            "summary_f1_precipitation_"
            f"{WINDOW_TAG}_all_rps.csv"
        ),
    )

    write_both(
        heat_comparison,
        METRICS_DIR
        / (
            "f1_comparison_vs_native_"
            f"heatindex_{WINDOW_TAG}_"
            "all_rps.csv"
        ),
    )

    write_both(
        precipitation_comparison,
        METRICS_DIR
        / (
            "f1_comparison_vs_native_"
            f"precipitation_{WINDOW_TAG}_"
            "all_rps.csv"
        ),
    )

    for metric in METRIC_NAMES:
        write_metric_workbook(
            confusion,
            metric,
        )

    write_count_workbook(
        confusion
    )

    write_audit_workbook(
        confusion
    )

    log(
        "Completed Stage 14 model "
        "evaluation post-processing"
    )


if __name__ == "__main__":
    main()
