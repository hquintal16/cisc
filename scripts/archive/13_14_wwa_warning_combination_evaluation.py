#!/usr/bin/env python3
"""Evaluate all selected Mesonet warning combinations against modeled hazards.

This replaces the former Stage 13 county evaluation and Stage 14 postprocessing
scripts. Warning combinations are evaluated independently for heat index and
precipitation over the fixed 2002-01-01 through 2023-12-31 window.

Shared, warning-independent products are built once under:
    data/out/13_performance/_shared/

Each warning combination is written under:
    data/out/13_performance/heat/<combination_tag>/
    data/out/13_performance/precipitation/<combination_tag>/

The final cross-combination ranking is written under:
    data/out/13_performance/warning_combination_summary/

Modes
-----
prepare_shared
    Build county/zone lookups, individual warning county-day files, modeled
    county-day files, and warning-combination manifests.

evaluate_worker
    Evaluate a chunk of warning combinations selected by SLURM_ARRAY_TASK_ID.
    Set CISC_COMBINATION_VARIABLE to heat or precipitation.

summarize_all
    Compile all completed combinations, rank F1 scores, identify the best
    warning combination for every dataset/return-period result, identify the
    overall best combination per variable, and write Excel workbooks.
"""

from __future__ import annotations

import itertools
import json
import math
import os
import re
import subprocess
import shutil
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Iterable, Sequence

import geopandas as gpd
import numpy as np
import pandas as pd
import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from scipy import stats
from shapely.geometry import box


# =============================================================================
# Configuration
# =============================================================================

ROOT = Path(os.environ.get("CISC_ROOT", "/proj/sebastianlab/users/hquintal/cisc"))
DATA_OUT = Path(os.environ.get("CISC_DATA_OUT", str(ROOT / "data/out")))
CLUSTER_BASE = Path(os.environ.get("CISC_CLUSTER_DIR", str(DATA_OUT / "06_clusters")))
PERF_BASE = Path(os.environ.get("CISC_PERF_DIR", str(DATA_OUT / "13_performance")))
WWA_DIR = Path(os.environ.get("CISC_WWA_DIR", str(ROOT / "figures/fig_02_wwa")))
ZONE_SHP = Path(
    os.environ.get(
        "CISC_ZONE_SHP",
        str(ROOT / "data/in/shapefile/nws_public_forecast_zone/z_18mr25.shp"),
    )
)
COUNTY_SHP_TEXT = os.environ.get("CISC_COUNTY_SHP", "").strip()
COUNTY_SHP = Path(COUNTY_SHP_TEXT) if COUNTY_SHP_TEXT else None

MODE = os.environ.get("CISC_WWA_COMBO_MODE", "prepare_shared").strip().lower()
COMBINATION_VARIABLE = os.environ.get("CISC_COMBINATION_VARIABLE", "heat").strip().lower()
WORKER_INDEX = int(os.environ.get("SLURM_ARRAY_TASK_ID", os.environ.get("CISC_WWA_WORKER_INDEX", "0")))
COMBINATIONS_PER_TASK = int(os.environ.get("CISC_COMBINATIONS_PER_TASK", "25"))

OVERWRITE_SHARED = os.environ.get("CISC_OVERWRITE_SHARED", "false").lower() in {"1", "true", "yes"}
OVERWRITE_COMBINATIONS = os.environ.get("CISC_OVERWRITE_COMBINATIONS", "false").lower() in {"1", "true", "yes"}
WRITE_ALL_EXCEL = os.environ.get("CISC_WRITE_ALL_EXCEL", "false").lower() in {"1", "true", "yes"}
WRITE_COUNTY_METRICS = os.environ.get("CISC_WRITE_COUNTY_METRICS", "false").lower() in {"1", "true", "yes"}

YEAR_START = 2002
YEAR_END = 2023
START = pd.Timestamp("2002-01-01")
END = pd.Timestamp("2023-12-31")
WINDOW_TAG = "20020101_20231231"
TOTAL_DAYS = int((END - START).days + 1)
BBOX = tuple(float(v) for v in os.environ.get("CISC_EVAL_BBOX", "-95,22,-75,40").split(","))

SCENARIOS = (
    "filtered_rp5",
    "filtered_rp10",
    "filtered_rp25",
    "filtered_rp50",
)
RETURN_PERIODS = (5, 10, 25, 50)
SCENARIO_BY_RP = {rp: f"filtered_rp{rp}" for rp in RETURN_PERIODS}
RP_BY_SCENARIO = {scenario: rp for rp, scenario in SCENARIO_BY_RP.items()}

# Warning products currently available in figures/fig_02_wwa.
# Keep these tuples synchronized only here; the SLURM driver reads their
# lengths directly from this Python file when sizing its arrays.
HEAT_WARNINGS = (
    "EH_W",
)

PRECIPITATION_WARNINGS = (
    # "HI_W",
    # "HU_W",
    # "TI_W",
    # "TR_W",
    "FF_W" #,
    # "SV_W",
    # "BZ_W",
    # "CF_W",
    # "FA_W",
    # "FL_W",
    # "HS_W",
    # "WS_W",
)

RESOLUTION_ORDER = (
    "Native",
    "STM1 bilinear",
    "STM1 max",
    "STM2 bilinear",
    "STM2 max",
)

DATASET_LABELS = {
    ("heatindex", "heatindex_0.25"): "Native",
    ("heatindex", "heatindex_0.30_bilinear"): "STM1 bilinear",
    ("heatindex", "heatindex_0.30_max"): "STM1 max",
    ("heatindex", "heatindex_0.36_bilinear"): "STM2 bilinear",
    ("heatindex", "heatindex_0.36_max"): "STM2 max",
    ("precipitation", "precipitation_0.25"): "Native",
    ("precipitation", "precipitation_0.11_bilinear"): "STM1 bilinear",
    ("precipitation", "precipitation_0.11_max"): "STM1 max",
    ("precipitation", "precipitation_0.14_bilinear"): "STM2 bilinear",
    ("precipitation", "precipitation_0.14_max"): "STM2 max",
}

NATIVE_DATASET = {
    "heatindex": "heatindex_0.25",
    "precipitation": "precipitation_0.25",
}

METRIC_NAMES = ("precision", "recall", "specificity", "accuracy", "f1", "f2")
COUNT_NAMES = ("TP", "FP", "FN", "TN", "total")

R_LIBS_USER = os.environ.get(
    "R_LIBS_USER",
    "/nas/longleaf/home/hquintal/R/x86_64-pc-linux-gnu-library/4.5",
)
RSCRIPT = os.environ.get("CISC_RSCRIPT", "Rscript")


# =============================================================================
# Paths
# =============================================================================

SHARED = PERF_BASE / "_shared"
CONTROL = SHARED / "control"
LOOKUP_XY = SHARED / "lookup_xy_to_county"
LOOKUP_ZONE = SHARED / "lookup_zone_to_county"
OBS_TAG_DIR = SHARED / "observed_warning_tags"
MODELED_DIR = SHARED / "modeled_county_daily"
TMP_DIR = SHARED / "tmp"

COUNTY_CACHE = CONTROL / "counties_2022_cb.gpkg"
COUNTY_ZIP_CACHE = CONTROL / "cb_2022_us_county_500k.zip"
ZONE_COUNTY = LOOKUP_ZONE / "lookup_zone_to_county_bbox.parquet"
DATASET_REGISTRY = CONTROL / "dataset_registry.csv"
HEAT_COMBINATION_MANIFEST = CONTROL / "heat_warning_combinations.csv"
PRECIP_COMBINATION_MANIFEST = CONTROL / "precipitation_warning_combinations.csv"
SUMMARY_DIR = PERF_BASE / "warning_combination_summary"

COUNTY_URL = os.environ.get(
    "CISC_COUNTY_URL",
    "https://www2.census.gov/geo/tiger/GENZ2022/shp/cb_2022_us_county_500k.zip",
)


# =============================================================================
# Generic helpers
# =============================================================================


def log(message: str) -> None:
    print(f"[{datetime.now():%Y-%m-%d %H:%M:%S}] {message}", flush=True)


def mkdirs() -> None:
    for path in (
        PERF_BASE,
        SHARED,
        CONTROL,
        LOOKUP_XY,
        LOOKUP_ZONE,
        OBS_TAG_DIR,
        MODELED_DIR,
        TMP_DIR,
        SUMMARY_DIR,
    ):
        path.mkdir(parents=True, exist_ok=True)


def atomic_csv(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp.csv")
    frame.to_csv(temporary, index=False)
    os.replace(temporary, path)


def atomic_parquet(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary_name = tempfile.mkstemp(
        prefix=f"{path.stem}.", suffix=".tmp.parquet", dir=path.parent
    )
    os.close(fd)
    temporary = Path(temporary_name)
    try:
        frame.to_parquet(
            temporary,
            index=False,
            engine="pyarrow",
            compression="zstd",
        )
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def write_both(frame: pd.DataFrame, csv_path: Path) -> None:
    atomic_csv(frame, csv_path)
    atomic_parquet(frame, csv_path.with_suffix(".parquet"))


def geoid5(values: Iterable) -> pd.Series:
    series = pd.Series(values, dtype="string").str.extract(r"(\d+)", expand=False)
    return series.where(series.str.len().le(5)).str.zfill(5)


def scenario_rp(name: str) -> int:
    match = re.fullmatch(r"filtered_rp(\d+)", name)
    if not match:
        raise ValueError(f"Unsupported scenario: {name}")
    return int(match.group(1))


def var_group(dataset_tag: str) -> str:
    tag = dataset_tag.lower()
    if tag.startswith("heatindex"):
        return "heatindex"
    if tag.startswith("precipitation"):
        return "precipitation"
    raise ValueError(f"Cannot infer variable from dataset tag: {dataset_tag}")


def parse_wwa_tag(path: Path) -> str | None:
    match = re.match(r"^wwa_freq_([A-Z]{2})_([A-Z])_.+\.qs$", path.name)
    return f"{match.group(1)}_{match.group(2)}".upper() if match else None


def combination_tag(warnings: Sequence[str]) -> str:
    return "__".join(warnings)


def experiment_dir(variable: str, combo_tag: str) -> Path:
    return PERF_BASE / variable / combo_tag


def manifest_path(variable: str) -> Path:
    if variable == "heat":
        return HEAT_COMBINATION_MANIFEST
    if variable == "precipitation":
        return PRECIP_COMBINATION_MANIFEST
    raise ValueError("Variable must be heat or precipitation")


def model_var(variable: str) -> str:
    return "heatindex" if variable == "heat" else "precipitation"


# =============================================================================
# Shared geospatial inputs
# =============================================================================


def download_county_shapefile() -> Path:
    if COUNTY_SHP is not None:
        if not COUNTY_SHP.exists():
            raise FileNotFoundError(f"Missing county shapefile: {COUNTY_SHP}")
        return COUNTY_SHP

    if COUNTY_CACHE.exists():
        return COUNTY_CACHE

    COUNTY_ZIP_CACHE.parent.mkdir(parents=True, exist_ok=True)
    log(f"Downloading county boundaries: {COUNTY_URL}")
    response = requests.get(COUNTY_URL, timeout=300)
    response.raise_for_status()
    temporary_zip = COUNTY_ZIP_CACHE.with_suffix(".tmp.zip")
    temporary_zip.write_bytes(response.content)
    os.replace(temporary_zip, COUNTY_ZIP_CACHE)

    extract_dir = CONTROL / "cb_2022_us_county_500k"
    extract_dir.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(COUNTY_ZIP_CACHE, "r") as archive:
        archive.extractall(extract_dir)

    shapefiles = sorted(extract_dir.glob("*.shp"))
    if not shapefiles:
        raise RuntimeError("County ZIP contained no shapefile")

    counties = gpd.read_file(shapefiles[0])
    counties.to_file(COUNTY_CACHE, driver="GPKG")
    return COUNTY_CACHE


def prepare_counties() -> gpd.GeoDataFrame:
    counties = gpd.read_file(download_county_shapefile())
    required = {"GEOID", "STATEFP"}
    missing = required.difference(counties.columns)
    if missing:
        raise ValueError(f"County data missing columns: {sorted(missing)}")

    counties = counties[
        ~counties["STATEFP"].astype(str).isin(
            {"02", "15", "60", "66", "69", "72", "78"}
        )
    ].copy()
    counties["GEOID"] = geoid5(counties["GEOID"])
    counties = counties[counties["GEOID"].notna()].to_crs(4326)
    counties = counties[counties.geometry.intersects(box(*BBOX))].copy()
    counties = counties[["GEOID", "geometry"]].copy()
    counties.geometry = counties.geometry.make_valid()
    return counties


def prepare_zones() -> gpd.GeoDataFrame:
    if not ZONE_SHP.exists():
        raise FileNotFoundError(ZONE_SHP)
    zones = gpd.read_file(ZONE_SHP)
    if "STATE_ZONE" not in zones.columns:
        raise ValueError("ZONE_SHP missing STATE_ZONE")
    zones["STATE_ZONE"] = zones["STATE_ZONE"].astype(str).str.strip().str.upper()
    zones = zones[zones["STATE_ZONE"].ne("")].to_crs(4326)
    zones.geometry = zones.geometry.make_valid()
    return zones[zones.geometry.intersects(box(*BBOX))][["STATE_ZONE", "geometry"]].copy()


def build_zone_county(counties: gpd.GeoDataFrame, zones: gpd.GeoDataFrame) -> pd.DataFrame:
    if ZONE_COUNTY.exists() and not OVERWRITE_SHARED:
        return pd.read_parquet(ZONE_COUNTY)

    joined = gpd.sjoin(
        zones.to_crs(5070), counties.to_crs(5070), how="inner", predicate="intersects"
    )
    output = (
        joined[["STATE_ZONE", "GEOID"]]
        .drop_duplicates()
        .sort_values(["STATE_ZONE", "GEOID"])
        .reset_index(drop=True)
    )
    atomic_parquet(output, ZONE_COUNTY)
    atomic_csv(output, ZONE_COUNTY.with_suffix(".csv"))
    return output


# =============================================================================
# Warning data
# =============================================================================


def qs_to_parquet(qs_path: Path, parquet_path: Path) -> Path:
    if parquet_path.exists() and not OVERWRITE_SHARED:
        return parquet_path

    resolved_rscript = shutil.which(RSCRIPT) if not Path(RSCRIPT).is_file() else RSCRIPT
    if not resolved_rscript:
        raise RuntimeError(
            "Rscript is unavailable. Run this workflow through the supplied "
            "SLURM driver, which loads R and installs/checks qs before Python starts."
        )

    parquet_path.parent.mkdir(parents=True, exist_ok=True)
    code = (
        "suppressPackageStartupMessages({library(qs);library(data.table);library(arrow)});"
        "x<-as.data.table(qread(commandArgs(TRUE)[1]));"
        "write_parquet(x,commandArgs(TRUE)[2],compression='zstd')"
    )
    environment = os.environ.copy()
    environment["R_LIBS_USER"] = R_LIBS_USER
    completed = subprocess.run(
        [str(resolved_rscript), "-e", code, str(qs_path), str(parquet_path)],
        env=environment,
        capture_output=True,
        text=True,
    )
    if completed.returncode:
        raise RuntimeError(
            f"Failed converting {qs_path}\n{completed.stdout}\n{completed.stderr}"
        )
    return parquet_path


def discover_warning_files() -> dict[str, Path]:
    discovered: dict[str, Path] = {}
    for path in sorted(WWA_DIR.glob("wwa_freq_*.qs")):
        tag = parse_wwa_tag(path)
        if tag:
            discovered[tag] = path

    required = set(HEAT_WARNINGS) | set(PRECIPITATION_WARNINGS)
    missing = sorted(required.difference(discovered))
    if missing:
        raise RuntimeError(
            "Missing required Mesonet warning files for: " + ", ".join(missing)
        )
    return discovered


def build_individual_warning_county_days(
    z2c: pd.DataFrame,
    warning_files: dict[str, Path],
) -> None:
    for index, warning_tag in enumerate(sorted(warning_files), start=1):
        if warning_tag not in set(HEAT_WARNINGS) | set(PRECIPITATION_WARNINGS):
            continue

        output_path = OBS_TAG_DIR / f"wwa_county_days_{warning_tag}.parquet"
        if output_path.exists() and not OVERWRITE_SHARED:
            continue

        qs_path = warning_files[warning_tag]
        cache_path = TMP_DIR / "wwa_qs_cache" / f"{qs_path.stem}.parquet"
        qs_to_parquet(qs_path, cache_path)
        frame = pd.read_parquet(cache_path)

        if not {"STATE_ZONE", "date"}.issubset(frame.columns):
            output = pd.DataFrame(columns=["warning_tag", "GEOID", "date"])
        else:
            frame = frame[["STATE_ZONE", "date"]].copy()
            frame["STATE_ZONE"] = frame["STATE_ZONE"].astype(str).str.strip().str.upper()
            frame["date"] = pd.to_datetime(frame["date"], errors="coerce").dt.normalize()
            frame = frame[frame["date"].between(START, END)]
            frame = frame.merge(z2c, on="STATE_ZONE", how="inner")
            output = (
                frame[["GEOID", "date"]]
                .drop_duplicates()
                .assign(warning_tag=warning_tag)
                [["warning_tag", "GEOID", "date"]]
                .sort_values(["date", "GEOID"])
                .reset_index(drop=True)
            )

        atomic_parquet(output, output_path)
        atomic_csv(output, output_path.with_suffix(".csv"))
        log(f"Warning tag {index}: wrote {warning_tag} ({len(output):,} county-days)")


# =============================================================================
# Modeled county-day cache
# =============================================================================


def discover_registry() -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for scenario in SCENARIOS:
        scenario_dir = CLUSTER_BASE / scenario
        if not scenario_dir.exists():
            raise FileNotFoundError(f"Missing scenario directory: {scenario_dir}")

        for dataset_dir in sorted(path for path in scenario_dir.iterdir() if path.is_dir()):
            files = sorted(dataset_dir.glob("cluster_rp*_*.parquet"))
            if files:
                rows.append(
                    {
                        "scenario": scenario,
                        "return_period": scenario_rp(scenario),
                        "dataset_tag": dataset_dir.name,
                        "var_group": var_group(dataset_dir.name),
                        "dataset_dir": str(dataset_dir),
                        "n_cluster_files": len(files),
                    }
                )

    registry = pd.DataFrame(rows)
    if registry.empty:
        raise RuntimeError(f"No cluster Parquet files found under {CLUSTER_BASE}")

    counts = registry.groupby(["scenario", "var_group"])["dataset_tag"].nunique()
    invalid = counts[counts.ne(5)]
    if not invalid.empty:
        raise RuntimeError(
            "Expected five datasets for every variable/return-period combination:\n"
            + invalid.to_string()
        )

    registry = registry.sort_values(
        ["var_group", "return_period", "dataset_tag"]
    ).reset_index(drop=True)
    atomic_csv(registry, DATASET_REGISTRY)
    return registry


def build_xy_lookup(dataset_tag: str, counties: gpd.GeoDataFrame) -> Path:
    output_path = LOOKUP_XY / f"lookup_xy_to_county_{dataset_tag}.parquet"
    if output_path.exists() and not OVERWRITE_SHARED:
        return output_path

    source_dir = CLUSTER_BASE / "clustered_rp5" / dataset_tag
    parts: list[pd.DataFrame] = []
    for path in sorted(source_dir.glob("cluster_rp5_*.parquet")):
        match = re.search(r"(\d{4})$", path.stem)
        if not match or not YEAR_START <= int(match.group(1)) <= YEAR_END:
            continue
        frame = pd.read_parquet(path, columns=["x", "y", "cluster"])
        frame["cluster"] = pd.to_numeric(frame["cluster"], errors="coerce")
        frame = frame[frame["cluster"].gt(0)][["x", "y"]].copy()
        frame = frame.apply(pd.to_numeric, errors="coerce").dropna().drop_duplicates()
        if not frame.empty:
            parts.append(frame)

    if not parts:
        raise RuntimeError(f"No positive RP5 coordinates found for {dataset_tag}")

    xy = pd.concat(parts, ignore_index=True).drop_duplicates().sort_values(["y", "x"])
    xy["x_key"] = xy["x"].round(8)
    xy["y_key"] = xy["y"].round(8)
    points = gpd.GeoDataFrame(
        xy, geometry=gpd.points_from_xy(xy["x"], xy["y"]), crs=4326
    )
    joined = gpd.sjoin(points, counties, how="left", predicate="within")

    missing = joined["GEOID"].isna()
    if missing.any():
        boundary = (
            gpd.sjoin(points.loc[missing], counties, how="left", predicate="intersects")
            .dropna(subset=["GEOID"])
            .reset_index()
            .drop_duplicates("index")
        )
        joined.loc[boundary["index"], "GEOID"] = boundary["GEOID"].to_numpy()

    output = (
        joined[["x", "y", "x_key", "y_key", "GEOID"]]
        .dropna(subset=["GEOID"])
        .drop_duplicates(["x_key", "y_key"])
        .sort_values(["y", "x"])
        .reset_index(drop=True)
    )
    output["GEOID"] = geoid5(output["GEOID"])
    atomic_parquet(output, output_path)
    atomic_csv(output, output_path.with_suffix(".csv"))
    return output_path


def modeled_output_path(scenario: str, dataset_tag: str) -> Path:
    return MODELED_DIR / scenario / f"modeled_county_daily_{dataset_tag}.parquet"


def build_modeled_county_days(row: pd.Series, counties: gpd.GeoDataFrame) -> Path:
    output_path = modeled_output_path(row["scenario"], row["dataset_tag"])
    if output_path.exists() and not OVERWRITE_SHARED:
        return output_path

    lookup_path = build_xy_lookup(row["dataset_tag"], counties)
    lookup = pd.read_parquet(lookup_path)[["x_key", "y_key", "GEOID"]].drop_duplicates()
    parts: list[pd.DataFrame] = []

    for path in sorted(Path(row["dataset_dir"]).glob("cluster_rp*_*.parquet")):
        match = re.search(r"(\d{4})$", path.stem)
        if not match or not YEAR_START <= int(match.group(1)) <= YEAR_END:
            continue

        frame = pd.read_parquet(path, columns=["x", "y", "time_utc", "cluster", "cluster_id"])
        frame["cluster"] = pd.to_numeric(frame["cluster"], errors="coerce")
        frame = frame[frame["cluster"].gt(0) & frame["cluster_id"].notna()].copy()
        if frame.empty:
            continue

        frame["date"] = (
            pd.to_datetime(frame["time_utc"], errors="coerce", utc=True)
            .dt.tz_convert(None)
            .dt.normalize()
        )
        frame = frame[frame["date"].between(START, END)]
        frame["x_key"] = pd.to_numeric(frame["x"], errors="coerce").round(8)
        frame["y_key"] = pd.to_numeric(frame["y"], errors="coerce").round(8)
        frame = frame.merge(lookup, on=["x_key", "y_key"], how="inner")
        if not frame.empty:
            parts.append(frame[["GEOID", "date"]].drop_duplicates())

    if parts:
        output = (
            pd.concat(parts, ignore_index=True)
            .drop_duplicates()
            .sort_values(["date", "GEOID"])
            .reset_index(drop=True)
        )
    else:
        output = pd.DataFrame(columns=["GEOID", "date"])

    atomic_parquet(output, output_path)
    return output_path


# =============================================================================
# Combination manifests
# =============================================================================


def all_nonempty_combinations(warnings: Sequence[str]) -> list[tuple[str, ...]]:
    output: list[tuple[str, ...]] = []
    for size in range(1, len(warnings) + 1):
        output.extend(itertools.combinations(warnings, size))
    return output


def write_combination_manifest(variable: str, warnings: Sequence[str]) -> pd.DataFrame:
    combinations = all_nonempty_combinations(warnings)
    rows = []
    for combination_index, selected in enumerate(combinations):
        tag = combination_tag(selected)
        rows.append(
            {
                "combination_index": combination_index,
                "variable": variable,
                "model_variable": model_var(variable),
                "combination_tag": tag,
                "n_warnings": len(selected),
                "warnings": ",".join(selected),
                "experiment_dir": str(experiment_dir(variable, tag)),
            }
        )

    manifest = pd.DataFrame(rows)
    atomic_csv(manifest, manifest_path(variable))
    log(f"Wrote {variable} combinations: {len(manifest):,}")
    return manifest


# =============================================================================
# Evaluation
# =============================================================================


def load_combined_observed(warnings: Sequence[str]) -> pd.DataFrame:
    parts: list[pd.DataFrame] = []
    for warning in warnings:
        path = OBS_TAG_DIR / f"wwa_county_days_{warning}.parquet"
        if not path.exists():
            raise FileNotFoundError(f"Missing observed warning cache: {path}")
        frame = pd.read_parquet(path, columns=["GEOID", "date"])
        if not frame.empty:
            parts.append(frame)

    if not parts:
        return pd.DataFrame(columns=["GEOID", "date"])

    output = pd.concat(parts, ignore_index=True).drop_duplicates()
    output["GEOID"] = geoid5(output["GEOID"])
    output["date"] = pd.to_datetime(output["date"], errors="coerce").dt.normalize()
    return output.dropna().drop_duplicates().sort_values(["GEOID", "date"])


def safe_divide(numerator: pd.Series, denominator: pd.Series) -> pd.Series:
    return numerator / denominator.replace(0, np.nan)


def evaluate_combination(manifest_row: pd.Series, registry: pd.DataFrame, counties: Sequence[str]) -> Path:
    variable = str(manifest_row["variable"])
    combo_tag = str(manifest_row["combination_tag"])
    selected_warnings = tuple(str(manifest_row["warnings"]).split(","))
    output_dir = experiment_dir(variable, combo_tag)
    metrics_path = output_dir / "metrics_county.parquet"
    confusion_path = output_dir / "confusion_matrices.parquet"
    completion_path = output_dir / "complete.json"

    if confusion_path.exists() and completion_path.exists() and not OVERWRITE_COMBINATIONS:
        return confusion_path

    output_dir.mkdir(parents=True, exist_ok=True)
    observed = load_combined_observed(selected_warnings)
    observed_counts = observed.groupby("GEOID").size().rename("obs_n")

    configuration = {
        "variable": variable,
        "model_variable": model_var(variable),
        "combination_tag": combo_tag,
        "warnings": list(selected_warnings),
        "n_warnings": len(selected_warnings),
        "window_tag": WINDOW_TAG,
        "start_date": str(START.date()),
        "end_date": str(END.date()),
        "total_days": TOTAL_DAYS,
        "scenarios": list(SCENARIOS),
    }
    (output_dir / "configuration.json").write_text(json.dumps(configuration, indent=2) + "\n")

    model_variable = model_var(variable)
    rows: list[pd.DataFrame] = []
    base = pd.DataFrame({"GEOID": pd.Series(counties, dtype="string")})

    for registry_row in registry[registry["var_group"].eq(model_variable)].itertuples(index=False):
        modeled_path = modeled_output_path(registry_row.scenario, registry_row.dataset_tag)
        if not modeled_path.exists():
            raise FileNotFoundError(f"Missing modeled cache: {modeled_path}")

        modeled = pd.read_parquet(modeled_path, columns=["GEOID", "date"])
        modeled["GEOID"] = geoid5(modeled["GEOID"])
        modeled["date"] = pd.to_datetime(modeled["date"], errors="coerce").dt.normalize()
        modeled = modeled.dropna().drop_duplicates()

        modeled_counts = modeled.groupby("GEOID").size().rename("mod_n")
        if observed.empty or modeled.empty:
            intersection_counts = pd.Series(dtype="int64", name="tp_n")
        else:
            intersection = observed.merge(modeled, on=["GEOID", "date"], how="inner")
            intersection_counts = intersection.groupby("GEOID").size().rename("tp_n")

        county = base.join(observed_counts, on="GEOID").join(modeled_counts, on="GEOID").join(intersection_counts, on="GEOID")
        county[["obs_n", "mod_n", "tp_n"]] = county[["obs_n", "mod_n", "tp_n"]].fillna(0).astype("int64")

        tp = county["tp_n"]
        fp = (county["mod_n"] - tp).clip(lower=0)
        fn = (county["obs_n"] - tp).clip(lower=0)
        tn = (TOTAL_DAYS - tp - fp - fn).clip(lower=0)
        total = tp + fp + fn + tn

        precision = safe_divide(tp, tp + fp)
        recall = safe_divide(tp, tp + fn)
        specificity = safe_divide(tn, tn + fp)
        accuracy = safe_divide(tp + tn, total)
        f1 = safe_divide(2 * tp, 2 * tp + fp + fn)
        f2 = safe_divide(5 * tp, 5 * tp + 4 * fn + fp)

        result = pd.DataFrame(
            {
                "variable": variable,
                "var_group": model_variable,
                "combination_tag": combo_tag,
                "warnings": ",".join(selected_warnings),
                "n_warnings": len(selected_warnings),
                "scenario": registry_row.scenario,
                "return_period": registry_row.return_period,
                "dataset_tag": registry_row.dataset_tag,
                "window_tag": WINDOW_TAG,
                "GEOID": county["GEOID"],
                "TP": tp,
                "FP": fp,
                "FN": fn,
                "TN": tn,
                "total": total,
                "precision": precision,
                "recall": recall,
                "specificity": specificity,
                "accuracy": accuracy,
                "f1": f1,
                "f2": f2,
            }
        )
        rows.append(result)

    metrics = pd.concat(rows, ignore_index=True)
    if WRITE_COUNTY_METRICS:
        atomic_parquet(metrics, metrics_path)
        atomic_csv(metrics, metrics_path.with_suffix(".csv"))

    confusion = build_confusion_summary(metrics)
    write_both(confusion, output_dir / "confusion_matrices.csv")

    county_f1 = summarize_county_f1(metrics)
    write_both(county_f1, output_dir / "summary_f1.csv")

    native_comparison = compare_f1_to_native(metrics, model_variable)
    write_both(native_comparison, output_dir / "f1_comparison_vs_native.csv")

    if WRITE_ALL_EXCEL:
        write_performance_workbooks(confusion, output_dir / "performance_metric_tables")

    completion_path.write_text(
        json.dumps(
            {
                "completed_at": datetime.now().isoformat(timespec="seconds"),
                "combination_tag": combo_tag,
                "warnings": list(selected_warnings),
                "n_metric_rows": len(metrics),
            },
            indent=2,
        )
        + "\n"
    )
    return confusion_path


def run_evaluate_worker() -> None:
    if COMBINATION_VARIABLE not in {"heat", "precipitation"}:
        raise ValueError("CISC_COMBINATION_VARIABLE must be heat or precipitation")
    if COMBINATIONS_PER_TASK < 1:
        raise ValueError("CISC_COMBINATIONS_PER_TASK must be at least 1")

    manifest = pd.read_csv(manifest_path(COMBINATION_VARIABLE))
    registry = pd.read_csv(DATASET_REGISTRY)
    counties = sorted(prepare_counties()["GEOID"].dropna().astype(str).unique())

    start_index = WORKER_INDEX * COMBINATIONS_PER_TASK
    stop_index = min(start_index + COMBINATIONS_PER_TASK, len(manifest))
    jobs = manifest.iloc[start_index:stop_index]
    log(
        f"{COMBINATION_VARIABLE} worker {WORKER_INDEX}: combinations "
        f"{start_index}-{max(start_index, stop_index - 1)} of {len(manifest):,}"
    )

    for position, (_, row) in enumerate(jobs.iterrows(), start=1):
        log(
            f"Evaluating {position}/{len(jobs)}: "
            f"{row['variable']} {row['combination_tag']}"
        )
        evaluate_combination(row, registry, counties)


# =============================================================================
# Postprocessing helpers
# =============================================================================


def derive_metrics(frame: pd.DataFrame) -> pd.DataFrame:
    output = frame.copy()
    output["total"] = output["TP"] + output["FP"] + output["FN"] + output["TN"]
    output["precision"] = safe_divide(output["TP"], output["TP"] + output["FP"])
    output["recall"] = safe_divide(output["TP"], output["TP"] + output["FN"])
    output["specificity"] = safe_divide(output["TN"], output["TN"] + output["FP"])
    output["accuracy"] = safe_divide(output["TP"] + output["TN"], output["total"])
    output["f1"] = safe_divide(2 * output["TP"], 2 * output["TP"] + output["FP"] + output["FN"])
    output["f2"] = safe_divide(5 * output["TP"], 5 * output["TP"] + 4 * output["FN"] + output["FP"])
    return output


def build_confusion_summary(metrics: pd.DataFrame) -> pd.DataFrame:
    grouped = (
        metrics.groupby(
            [
                "variable",
                "combination_tag",
                "warnings",
                "n_warnings",
                "var_group",
                "dataset_tag",
                "return_period",
            ],
            dropna=False,
        )
        .agg(TP=("TP", "sum"), FP=("FP", "sum"), FN=("FN", "sum"), TN=("TN", "sum"), N_rows=("GEOID", "size"))
        .reset_index()
    )
    grouped = derive_metrics(grouped)
    grouped["resolution"] = [
        DATASET_LABELS.get((variable, dataset_tag), pd.NA)
        for variable, dataset_tag in zip(grouped["var_group"], grouped["dataset_tag"])
    ]
    grouped = grouped[grouped["resolution"].notna()].copy()
    grouped["_resolution_order"] = grouped["resolution"].map(
        {label: index for index, label in enumerate(RESOLUTION_ORDER)}
    )
    grouped = grouped.sort_values(["return_period", "_resolution_order"]).drop(columns="_resolution_order")
    return grouped.reset_index(drop=True)


def summarize_county_f1(metrics: pd.DataFrame) -> pd.DataFrame:
    return (
        metrics.groupby(
            [
                "variable",
                "combination_tag",
                "warnings",
                "n_warnings",
                "scenario",
                "dataset_tag",
                "return_period",
                "window_tag",
            ],
            dropna=False,
        )["f1"]
        .agg(n_counties="count", mean_f1="mean", median_f1="median", sd_f1="std", min_f1="min", max_f1="max")
        .reset_index()
        .sort_values(["return_period", "dataset_tag"])
        .reset_index(drop=True)
    )


def paired_test_results(dataset_values: np.ndarray, native_values: np.ndarray) -> dict[str, float]:
    valid = np.isfinite(dataset_values) & np.isfinite(native_values)
    dataset_values = dataset_values[valid]
    native_values = native_values[valid]
    differences = dataset_values - native_values

    if differences.size == 0:
        return {key: np.nan for key in ("diff_mean", "ci_low", "ci_high", "t_p_value", "wilcox_p_value")}

    diff_mean = float(np.mean(differences))
    if differences.size >= 2:
        t_result = stats.ttest_rel(dataset_values, native_values, nan_policy="omit")
        standard_error = stats.sem(differences, nan_policy="omit")
        t_p_value = float(t_result.pvalue)
        if np.isfinite(standard_error):
            margin = 0.0 if standard_error == 0 else float(stats.t.ppf(0.975, differences.size - 1) * standard_error)
            ci_low, ci_high = diff_mean - margin, diff_mean + margin
        else:
            ci_low = ci_high = np.nan
    else:
        t_p_value = ci_low = ci_high = np.nan

    if differences.size and np.any(differences != 0):
        try:
            wilcox_p_value = float(stats.wilcoxon(dataset_values, native_values).pvalue)
        except ValueError:
            wilcox_p_value = np.nan
    else:
        wilcox_p_value = np.nan

    return {
        "diff_mean": diff_mean,
        "ci_low": ci_low,
        "ci_high": ci_high,
        "t_p_value": t_p_value,
        "wilcox_p_value": wilcox_p_value,
    }


def compare_f1_to_native(metrics: pd.DataFrame, variable: str) -> pd.DataFrame:
    native_tag = NATIVE_DATASET[variable]
    rows: list[dict[str, object]] = []

    for return_period in RETURN_PERIODS:
        scenario_data = metrics[metrics["return_period"].eq(return_period)].copy()
        native = (
            scenario_data.loc[
                scenario_data["dataset_tag"].eq(native_tag) & scenario_data["f1"].notna(),
                ["GEOID", "f1"],
            ]
            .rename(columns={"f1": "native_f1"})
            .drop_duplicates("GEOID")
        )

        for dataset_tag in sorted(set(scenario_data["dataset_tag"].dropna().astype(str)) - {native_tag}):
            dataset = (
                scenario_data.loc[
                    scenario_data["dataset_tag"].eq(dataset_tag) & scenario_data["f1"].notna(),
                    ["GEOID", "f1"],
                ]
                .rename(columns={"f1": "dataset_f1"})
                .drop_duplicates("GEOID")
            )
            paired = native.merge(dataset, on="GEOID", how="inner")
            tests = paired_test_results(
                paired["dataset_f1"].to_numpy(dtype=float),
                paired["native_f1"].to_numpy(dtype=float),
            ) if not paired.empty else {key: np.nan for key in ("diff_mean", "ci_low", "ci_high", "t_p_value", "wilcox_p_value")}

            rows.append(
                {
                    "variable": variable,
                    "return_period": return_period,
                    "dataset_tag": dataset_tag,
                    "label": DATASET_LABELS.get((variable, dataset_tag), dataset_tag),
                    "mean_dataset": paired["dataset_f1"].mean() if not paired.empty else np.nan,
                    "mean_native": paired["native_f1"].mean() if not paired.empty else np.nan,
                    **tests,
                    "n_paired": len(paired),
                }
            )

    return pd.DataFrame(rows)


# =============================================================================
# Excel output
# =============================================================================


def metric_matrix(confusion: pd.DataFrame, variable: str, metric: str) -> pd.DataFrame:
    subset = confusion[["resolution", "return_period", metric]].copy()
    template = pd.MultiIndex.from_product(
        [RESOLUTION_ORDER, RETURN_PERIODS], names=["resolution", "return_period"]
    ).to_frame(index=False)
    subset = template.merge(subset, on=["resolution", "return_period"], how="left")
    wide = subset.pivot(index="resolution", columns="return_period", values=metric)
    wide = wide.reindex(index=RESOLUTION_ORDER, columns=RETURN_PERIODS)
    wide.columns = [f"{rp}-year RP" for rp in wide.columns]
    return wide.reset_index()


def style_sheet(worksheet, title: str, table: pd.DataFrame, metric: bool) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet["A1"] = title
    worksheet["A1"].font = Font(name="Calibri", size=12, bold=True)
    header_row = 3
    bottom = Side(style="thin", color="000000")

    for column_index, column_name in enumerate(table.columns, start=1):
        cell = worksheet.cell(row=header_row, column=column_index, value=column_name)
        cell.font = Font(name="Calibri", size=10, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=bottom)

    for row_index, values in enumerate(table.itertuples(index=False, name=None), start=header_row + 1):
        for column_index, value in enumerate(values, start=1):
            cell = worksheet.cell(row=row_index, column=column_index, value=None if pd.isna(value) else value)
            cell.font = Font(name="Calibri", size=10, bold=column_index == 1)
            cell.alignment = Alignment(horizontal="left" if column_index == 1 else "center")
            if column_index > 1 and value is not None and not pd.isna(value):
                cell.number_format = "0.0000" if metric else "#,##0"

    worksheet.column_dimensions["A"].width = 19
    for column_index in range(2, len(table.columns) + 1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = 14
    worksheet.freeze_panes = "B4"


def write_performance_workbooks(confusion: pd.DataFrame, table_dir: Path) -> None:
    table_dir.mkdir(parents=True, exist_ok=True)
    variable = str(confusion["var_group"].iloc[0])

    for metric_name in METRIC_NAMES:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = variable
        table = metric_matrix(confusion, variable, metric_name)
        style_sheet(worksheet, f"{variable}: {metric_name} by resolution and return period", table, metric=True)
        workbook.save(table_dir / f"{metric_name}.xlsx")

    workbook = Workbook()
    workbook.remove(workbook.active)
    for count_name in COUNT_NAMES:
        worksheet = workbook.create_sheet(title=count_name)
        table = metric_matrix(confusion, variable, count_name)
        style_sheet(worksheet, f"{variable}: {count_name} by resolution and return period", table, metric=False)
    workbook.save(table_dir / "confusion_matrix_counts.xlsx")

    audit_columns = [
        "var_group", "combination_tag", "warnings", "dataset_tag", "resolution",
        "return_period", "TP", "FP", "FN", "TN", "total", *METRIC_NAMES,
    ]
    audit = confusion[audit_columns].copy()
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = variable
    for column_index, column_name in enumerate(audit.columns, start=1):
        worksheet.cell(row=1, column=column_index, value=column_name).font = Font(bold=True)
    for row_index, values in enumerate(audit.itertuples(index=False, name=None), start=2):
        for column_index, value in enumerate(values, start=1):
            worksheet.cell(row=row_index, column=column_index, value=None if pd.isna(value) else value)
    worksheet.freeze_panes = "A2"
    workbook.save(table_dir / "all_performance_metrics.xlsx")


# =============================================================================
# Cross-combination summary
# =============================================================================


def read_completed_confusions(variable: str) -> pd.DataFrame:
    manifest = pd.read_csv(manifest_path(variable))
    parts: list[pd.DataFrame] = []
    missing: list[str] = []

    for row in manifest.itertuples(index=False):
        path = experiment_dir(variable, row.combination_tag) / "confusion_matrices.parquet"
        if path.exists():
            parts.append(pd.read_parquet(path))
        else:
            missing.append(row.combination_tag)

    if missing:
        raise RuntimeError(
            f"{variable}: {len(missing):,} combinations are incomplete; first: {missing[:10]}"
        )
    return pd.concat(parts, ignore_index=True)


def rank_combinations(confusion: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    ranking = confusion.copy()
    ranking["rank_within_result"] = (
        ranking.groupby(["var_group", "dataset_tag", "return_period"])["f1"]
        .rank(method="min", ascending=False, na_option="bottom")
        .astype("Int64")
    )
    ranking = ranking.sort_values(
        ["var_group", "dataset_tag", "return_period", "rank_within_result", "n_warnings", "combination_tag"]
    ).reset_index(drop=True)

    best_by_result = ranking[ranking["rank_within_result"].eq(1)].copy()
    best_by_result = best_by_result.sort_values(
        ["var_group", "return_period", "resolution", "n_warnings", "combination_tag"]
    ).reset_index(drop=True)

    overall = (
        confusion.groupby(
            ["variable", "var_group", "combination_tag", "warnings", "n_warnings"],
            dropna=False,
        )
        .agg(
            mean_f1_across_results=("f1", "mean"),
            median_f1_across_results=("f1", "median"),
            min_f1_across_results=("f1", "min"),
            max_f1_across_results=("f1", "max"),
            n_results=("f1", "count"),
        )
        .reset_index()
    )
    overall["overall_rank"] = (
        overall.groupby("var_group")["mean_f1_across_results"]
        .rank(method="min", ascending=False, na_option="bottom")
        .astype("Int64")
    )
    overall = overall.sort_values(
        ["var_group", "overall_rank", "n_warnings", "combination_tag"]
    ).reset_index(drop=True)
    overall_best = overall[overall["overall_rank"].eq(1)].copy()
    return ranking, best_by_result, overall_best


def write_summary_workbook(
    ranking: pd.DataFrame,
    best_by_result: pd.DataFrame,
    overall_best: pd.DataFrame,
    output_path: Path,
) -> None:
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        overall_best.to_excel(writer, sheet_name="overall_best", index=False)
        best_by_result.to_excel(writer, sheet_name="best_by_result", index=False)
        ranking.to_excel(writer, sheet_name="all_rankings", index=False)

        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.sheet_view.showGridLines = False
            for cell in worksheet[1]:
                cell.font = Font(name="Calibri", size=10, bold=True)
            for column_index, column_cells in enumerate(worksheet.columns, start=1):
                maximum = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
                worksheet.column_dimensions[get_column_letter(column_index)].width = min(max(maximum + 2, 12), 45)


def run_summarize_all() -> None:
    heat = read_completed_confusions("heat")
    precipitation = read_completed_confusions("precipitation")
    confusion = pd.concat([heat, precipitation], ignore_index=True)

    ranking, best_by_result, overall_best = rank_combinations(confusion)
    write_both(confusion, SUMMARY_DIR / "all_warning_combination_confusion_matrices.csv")
    write_both(ranking, SUMMARY_DIR / "all_warning_combination_f1_rankings.csv")
    write_both(best_by_result, SUMMARY_DIR / "best_warning_combinations_by_result.csv")
    write_both(overall_best, SUMMARY_DIR / "overall_best_warning_combination_by_variable.csv")

    write_summary_workbook(
        ranking,
        best_by_result,
        overall_best,
        SUMMARY_DIR / "warning_combination_f1_summary.xlsx",
    )

    for row in overall_best.itertuples(index=False):
        source_dir = experiment_dir(row.variable, row.combination_tag)
        confusion_path = source_dir / "confusion_matrices.parquet"
        best_confusion = pd.read_parquet(confusion_path)
        output_dir = SUMMARY_DIR / "best_performance_metric_tables" / row.variable
        write_performance_workbooks(best_confusion, output_dir)

    log("Overall best combinations:")
    for row in overall_best.itertuples(index=False):
        log(
            f"  {row.var_group}: {row.warnings}; "
            f"mean F1 across 20 results={row.mean_f1_across_results:.6f}"
        )


# =============================================================================
# Workflow modes
# =============================================================================


def run_prepare_shared() -> None:
    counties = prepare_counties()
    zones = prepare_zones()
    zone_county = build_zone_county(counties, zones)
    warning_files = discover_warning_files()
    build_individual_warning_county_days(zone_county, warning_files)

    registry = discover_registry()
    unique_datasets = registry[["dataset_tag"]].drop_duplicates()
    for index, dataset_tag in enumerate(unique_datasets["dataset_tag"], start=1):
        log(f"XY lookup {index}/{len(unique_datasets)}: {dataset_tag}")
        build_xy_lookup(dataset_tag, counties)

    for index, (_, row) in enumerate(registry.iterrows(), start=1):
        log(
            f"Modeled cache {index}/{len(registry)}: "
            f"{row['scenario']} {row['dataset_tag']}"
        )
        build_modeled_county_days(row, counties)

    heat_manifest = write_combination_manifest("heat", HEAT_WARNINGS)
    precipitation_manifest = write_combination_manifest(
        "precipitation", PRECIPITATION_WARNINGS
    )

    run_information = {
        "heat_combinations": len(heat_manifest),
        "precipitation_combinations": len(precipitation_manifest),
        "combinations_per_task": COMBINATIONS_PER_TASK,
        "heat_array_tasks": math.ceil(len(heat_manifest) / COMBINATIONS_PER_TASK),
        "precipitation_array_tasks": math.ceil(len(precipitation_manifest) / COMBINATIONS_PER_TASK),
    }
    (CONTROL / "run_information.json").write_text(json.dumps(run_information, indent=2) + "\n")
    log(json.dumps(run_information))


def main() -> None:
    mkdirs()
    log(f"MODE={MODE} ROOT={ROOT}")
    if MODE == "prepare_shared":
        run_prepare_shared()
    elif MODE == "evaluate_worker":
        run_evaluate_worker()
    elif MODE == "summarize_all":
        run_summarize_all()
    else:
        raise ValueError("CISC_WWA_COMBO_MODE must be prepare_shared, evaluate_worker, or summarize_all")
    log("Done")


if __name__ == "__main__":
    main()
